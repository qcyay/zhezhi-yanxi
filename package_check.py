"""Offline verification of bundled parsers, OCR models and static resources."""
import json
import os
from pathlib import Path
import tempfile


def check():
    import app
    import formats
    import xlrd
    import onnxruntime
    from openpyxl import Workbook, load_workbook
    from pypdf import PdfReader, PdfWriter
    from PIL import Image, ImageDraw, ImageFont

    for name in ('static/index.html', 'static/app.js', 'reference_sources.json',
                 'starter_recommendations.json', 'VERSION'):
        assert (app.ROOT / name).is_file(), name
    for name in ('reference_sources.json', 'starter_recommendations.json'):
        json.loads((app.ROOT / name).read_text('utf-8'))
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        book = Workbook()
        book.active['A1'] = 'EXE test'
        book.save(root / 'sample.xlsx')
        loaded = load_workbook(root / 'sample.xlsx')
        assert loaded.active['A1'].value == 'EXE test'
        loaded.close()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(root / 'sample.pdf')
        assert len(PdfReader(root / 'sample.pdf').pages) == 1
        picture = Image.new('RGB', (700, 160), 'white')
        font = ImageFont.truetype(str(Path(os.environ['WINDIR']) / 'Fonts' / 'arial.ttf'), 60)
        ImageDraw.Draw(picture).text((20, 35), 'TEST 123', font=font, fill='black')
        picture.save(root / 'sample.png')
        text = formats.image_text(root / 'sample.png')
        assert '123' in text, repr(text)
    print('PASS: bundled PDF, XLSX, XLS imports, OCR inference and application resources', flush=True)
