"""Local image OCR / legacy Excel reading and content-addressed extraction cache."""
from pathlib import Path
import hashlib,sys,threading
VENDOR=Path(__file__).parent/'.deps'
if VENDOR.is_dir() and str(VENDOR) not in sys.path:sys.path.insert(0,str(VENDOR))
OCR_LOCK=threading.Lock();OCR_ENGINE=None

def image_text(path):
 global OCR_ENGINE
 from PIL import Image,ImageOps
 import numpy as np
 from rapidocr_onnxruntime import RapidOCR
 with Image.open(path) as raw:
  if raw.width*raw.height>40_000_000:raise ValueError('图片尺寸过大，请缩小后重新读取')
  picture=ImageOps.exif_transpose(raw).convert('RGB');picture.thumbnail((3200,3200));pixels=np.asarray(picture)[:,:,::-1].copy()
 with OCR_LOCK:
  if OCR_ENGINE is None:OCR_ENGINE=RapidOCR(intra_op_num_threads=2,inter_op_num_threads=1)
  result,_=OCR_ENGINE(pixels)
 text='\n'.join(str(row[1]).strip() for row in (result or []) if float(row[2])>=.45)
 if not text:raise ValueError('图片中未识别到清晰文字，请检查图片清晰度')
 return text

def xls_text(path):
 import xlrd
 from openpyxl.utils import get_column_letter
 book=xlrd.open_workbook(str(path),on_demand=True);lines=[];cells=0
 try:
  for sheet in book.sheets():
   if sheet.nrows>20000:raise ValueError('工作表行数较多，请拆分后读取')
   for row in range(sheet.nrows):
    values=[]
    for column in range(sheet.ncols):
     cells+=1
     if cells>200000:raise ValueError('表格内容较多，请拆分后读取')
     cell=sheet.cell(row,column)
     if cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK):continue
     value=cell.value
     if cell.ctype==xlrd.XL_CELL_DATE:value=xlrd.xldate_as_datetime(value,book.datemode).isoformat(sep=' ')
     elif cell.ctype==xlrd.XL_CELL_NUMBER and float(value).is_integer():value=int(value)
     elif cell.ctype==xlrd.XL_CELL_BOOLEAN:value='是' if value else '否'
     values.append(get_column_letter(column+1)+'='+str(value))
    if values:lines.append(f'工作表【{sheet.name}】第{row+1}行：'+'；'.join(values))
  return '\n'.join(lines)
 finally:book.release_resources()

def cached_extract(path,fallback,cache_dir):
 if path.stat().st_size>15_000_000:raise ValueError('文件较大，请拆分后重新读取')
 digest=hashlib.sha256(path.read_bytes()).hexdigest();key='v1-'+path.suffix.lower().lstrip('.')+'-'+digest;target=cache_dir/(key+'.txt')
 if target.is_file():return target.read_text('utf-8')
 suffix=path.suffix.lower()
 text=image_text(path) if suffix in ('.jpg','.jpeg','.png','.bmp','.webp','.tif','.tiff') else xls_text(path) if suffix=='.xls' else fallback(path)
 if text.strip() and len(text)<=600000:
  cache_dir.mkdir(parents=True,exist_ok=True);temp=target.with_suffix('.'+str(threading.get_ident())+'.tmp');temp.write_text(text,'utf-8');temp.replace(target)
 return text
